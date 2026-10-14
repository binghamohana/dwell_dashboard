"""Builds the Dwell Central setup workbook.

Usage: python3 tools/make_template.py public/assets/templates/Dwell-Central-template.xlsx [prefill.json]
Requires Python 3 and openpyxl. The optional JSON pre-fills the sheets (never commit a pre-filled copy:
it contains personnel data)."""
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]
PREFILL = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}

BRAND = '113263'
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
HEAD_FONT = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
NOTE_FILL = PatternFill('solid', fgColor='F4F1EC')
NOTE_FONT = Font(italic=True, color='6B737B')
CHECK_FONT = Font(color='9A5B2E')
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']

wb = Workbook()


def sheet(title, columns, rows, max_rows):
    """columns: [(header, width, note)]; rows: list of lists."""
    ws = wb.create_sheet(title)
    for i, (head, width, note) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=head)
        c.fill = NOTE_FILL if '(not imported)' in head or '(automatic)' in head else HEAD_FILL
        c.font = Font(bold=True, color='6B737B') if c.fill is NOTE_FILL else HEAD_FONT
        c.alignment = Alignment(vertical='center', wrap_text=True)
        if note:
            c.comment = Comment(note, 'Dwell Central', width=320, height=140)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'B2'
    for r, values in enumerate(rows, start=2):
        for i, v in enumerate(values, start=1):
            if v is None or v == '':
                continue
            if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-':
                v = date.fromisoformat(v)
            ws.cell(row=r, column=i, value=v)
    for i, (head, _, _) in enumerate(columns, start=1):
        if '(not imported)' in head:
            for r in range(2, max_rows + 1):
                ws.cell(row=r, column=i).font = CHECK_FONT
                ws.cell(row=r, column=i).alignment = Alignment(wrap_text=True, vertical='top')
    return ws


def dropdown(ws, col, formula, max_rows, strict=True, prompt=None):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
    dv.showErrorMessage = strict
    if strict:
        dv.error = 'Please pick from the list.'
        dv.errorTitle = 'Dwell Central'
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f'{col}2:{col}{max_rows}')


# ------------------------------------------------------------------ Start here
ws = wb.active
ws.title = 'Start here'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 110
lines = [
    ('Dwell Central setup workbook', Font(bold=True, size=18, color=BRAND)),
    ('Fill in the sheets in order, then upload this file in Dwell Central → Import. Nothing is saved until you review it and click Apply.', Font(size=11)),
    ('', None),
    ('1.  Companies — the brands people belong to (Dwell Communities, 58 Place). Already filled in; add more only if needed.', None),
    ('2.  Legal entities — each employer of record (the company on the paycheck). Type: Employer, or Contractors (1099) for independent contractors and vendors. Company: which brand it belongs to.', None),
    ('3.  Properties — every apartment community, office and team (for example Corporate Office, Landscape Crew). Property manager: the person who runs it; they are reminded about celebrations when someone\'s supervisor isn\'t available.', None),
    ('4.  Departments — for example Leasing, Maintenance, Property Management, Accounting, Leadership.', None),
    ('5.  People — one row per current team member. Department, Legal entity, Property / office, Supervisor and Dotted-line manager are drop-down lists that come from the other sheets, so everything links up.', None),
    ('6.  Upload: Dwell Central → Import → choose this file → review every change → Apply.', None),
    ('', None),
    ('Good to know', Font(bold=True, size=12, color=BRAND)),
    ('•  Supervisor: pick the person from the list (or type their work email). Supervisors decide who gets each birthday and anniversary reminder, so this matters most.', None),
    ('•  Birthday: month and day only. The year is not needed and not kept.', None),
    ('•  Hire date: the most recent start date. Original hire date: only if they worked here before and were rehired.', None),
    ('•  Status: Active, On leave, or Former employee. Former employees are only kept if they are already in Dwell Central.', None),
    ('•  "Also known as": other spellings used in HR or payroll exports (one per line, or separated by ;), so future HR files match automatically.', None),
    ('•  Columns with a grey heading are notes for you and are not imported. Please don\'t rename the sheets or the column headings.', None),
    ('•  You can upload an updated copy any time. Unchanged rows are skipped, blank cells never erase saved information, and records marked Verified in Dwell Central are protected.', None),
    ('•  Once filled in, this file contains personal information: keep it somewhere private and delete stray copies.', None),
]
for i, (text, font) in enumerate(lines, start=2):
    c = ws.cell(row=i, column=2, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if font:
        c.font = font
    ws.row_dimensions[i].height = 22 if len(text) < 110 else 36
if PREFILL.get('intro'):
    c = ws.cell(row=len(lines) + 3, column=2, value=PREFILL['intro'])
    c.font = Font(bold=True, color='9A5B2E')
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[len(lines) + 3].height = 60

# ------------------------------------------------------------------ organization sheets
ORG_ROWS = 200
PEOPLE_ROWS = 500
companies = PREFILL.get('companies', [['Dwell Communities', '#113263'], ['58 Place', '#8A4B6E']])
sheet('Companies', [
    ('Company', 30, 'The brand people belong to.'),
    ('Color', 12, 'Optional. Used in the org chart, like #113263.'),
    ('Notes (not imported)', 50, None),
], companies, ORG_ROWS)

ws = sheet('Legal entities', [
    ('Legal entity', 34, 'Short name used in Dwell Central, e.g. "Dwell Communities, LLC".'),
    ('Full legal name', 34, 'Optional.'),
    ('Type', 20, 'Employer, or Contractors (1099) for independent contractors and vendors.'),
    ('Company', 22, 'Which brand (from the Companies sheet).'),
    ('Also known as', 34, 'Other spellings in HR or payroll files, one per line or separated by ;'),
    ('Notes (not imported)', 50, None),
], PREFILL.get('entities', []), ORG_ROWS)
dropdown(ws, 'C', '"Employer,Contractors (1099),Other"', ORG_ROWS)
dropdown(ws, 'D', "=Companies!$A$2:$A$200", ORG_ROWS)

ws = sheet('Properties', [
    ('Property / office', 30, 'The name everyone uses, e.g. "Maple Commons" or "Corporate Office".'),
    ('Short name', 14, 'Optional.'),
    ('Type', 14, 'Property, Office, Team (e.g. a landscape crew) or Other.'),
    ('Company', 22, 'Which brand (from the Companies sheet).'),
    ('Address', 34, 'Optional.'),
    ('Property manager', 28, 'Optional. Pick from the list once the People sheet is filled in (or type their work email).'),
    ('Also known as', 28, 'Other spellings in HR files (e.g. the "Division" column), one per line or separated by ;'),
    ('Notes (not imported)', 50, None),
], PREFILL.get('properties', []), ORG_ROWS)
dropdown(ws, 'C', '"Property,Office,Team,Other"', ORG_ROWS)
dropdown(ws, 'D', "=Companies!$A$2:$A$200", ORG_ROWS)
dropdown(ws, 'F', "=People!$T$2:$T$500", ORG_ROWS, strict=False)

sheet('Departments', [
    ('Department', 30, 'e.g. Leasing, Maintenance, Property Management, Accounting, Leadership.'),
    ('Also known as', 30, 'Optional other spellings, separated by ;'),
    ('Notes (not imported)', 50, None),
], PREFILL.get('departments', []), ORG_ROWS)

# ------------------------------------------------------------------ people
people_cols = [
    ('Employee ID', 13, 'Optional: the ID from your HR / payroll system. Makes future imports match exactly.'),
    ('First name', 16, 'Legal first name.'),
    ('Last name', 20, 'Legal last name (all surnames).'),
    ('Goes by', 12, 'Optional: the first name they use, if different.'),
    ('Work email', 30, 'Needed for anyone who supervises others or will sign in.'),
    ('Job title', 26, None),
    ('Department', 22, 'From the Departments sheet.'),
    ('Legal entity', 28, 'From the Legal entities sheet.'),
    ('Property / office', 24, 'Where they mainly work (from the Properties sheet).'),
    ('Supervisor', 24, 'Pick from the list (or type a work email). Leave blank only for the top of the organization.'),
    ('Dotted-line manager', 22, 'Optional second manager.'),
    ('Status', 16, 'Active, On leave or Former employee.'),
    ('Classification', 20, None),
    ('Hire date', 13, 'Most recent start date.'),
    ('Original hire date', 13, 'Only if they were rehired: their first start date.'),
    ('Term date', 13, 'Only for former employees.'),
    ('Birth month', 13, 'Month only (no year needed).'),
    ('Birth day', 10, 'Day of the month, 1–31.'),
    ('To check (not imported)', 48, None),
    ('Lookup name (automatic)', 24, 'Filled in automatically; used by the Supervisor lists. Leave as is.'),
]
rows = PREFILL.get('people', [])
ws = sheet('People', people_cols, rows, PEOPLE_ROWS)
for r in range(2, PEOPLE_ROWS + 1):
    c = ws.cell(row=r, column=20, value=f'=IF(B{r}="","",TRIM(B{r}&" "&C{r}))')
    c.font = Font(color='9AA0A6')
    for col in (14, 15, 16):
        ws.cell(row=r, column=col).number_format = 'm/d/yyyy'
ws.column_dimensions['T'].hidden = True
dropdown(ws, 'G', "=Departments!$A$2:$A$200", PEOPLE_ROWS)
dropdown(ws, 'H', "='Legal entities'!$A$2:$A$200", PEOPLE_ROWS)
dropdown(ws, 'I', "=Properties!$A$2:$A$200", PEOPLE_ROWS)
dropdown(ws, 'J', "=$T$2:$T$500", PEOPLE_ROWS, strict=False)
dropdown(ws, 'K', "=$T$2:$T$500", PEOPLE_ROWS, strict=False)
dropdown(ws, 'L', '"Active,On leave,Former employee"', PEOPLE_ROWS)
dropdown(ws, 'M', '"Full time,Part time,Contractor / vendor,Temporary,Other"', PEOPLE_ROWS)
dropdown(ws, 'Q', '"' + ','.join(MONTHS) + '"', PEOPLE_ROWS)
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='31', allow_blank=True)
dv.error = 'Enter a day from 1 to 31.'
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add(f'R2:R{PEOPLE_ROWS}')
for col in ('N', 'O', 'P'):
    dvd = DataValidation(type='date', operator='greaterThan', formula1='1', allow_blank=True)
    dvd.error = 'Enter a date, e.g. 3/15/2021.'
    dvd.showErrorMessage = True
    ws.add_data_validation(dvd)
    dvd.add(f'{col}2:{col}{PEOPLE_ROWS}')

wb.active = 0
wb.save(OUT)
print('saved', OUT, len(rows), 'people')
